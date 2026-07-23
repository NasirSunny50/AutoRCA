from __future__ import annotations

import json
import logging
import os
import sqlite3
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import requests
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "dashboard.db"
LOG_PATH = BASE_DIR / "application.log"

app = Flask(__name__)
app.config["JSON_SORT_KEYS"] = False

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"), logging.StreamHandler()],
)
logger = logging.getLogger("enterprise-health-dashboard")


def db_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with db_connection() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS projects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                color TEXT NOT NULL DEFAULT '#0d6efd',
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS components (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                health_url TEXT NOT NULL,
                info_url TEXT,
                metrics_url TEXT,
                auth_type TEXT NOT NULL DEFAULT 'none',
                username TEXT,
                password TEXT,
                bearer_token TEXT,
                custom_headers TEXT,
                timeout INTEGER NOT NULL DEFAULT 5,
                enabled INTEGER NOT NULL DEFAULT 1,
                status TEXT NOT NULL DEFAULT 'UNKNOWN',
                http_status INTEGER,
                response_time REAL,
                version TEXT,
                build_date TEXT,
                git_branch TEXT,
                git_commit TEXT,
                java_version TEXT,
                cpu_usage TEXT,
                memory_usage TEXT,
                disk_usage TEXT,
                last_checked TEXT,
                error_message TEXT,
                raw_response TEXT,
                created_at TEXT NOT NULL,
                FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS health_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                component_id INTEGER NOT NULL,
                status TEXT NOT NULL,
                http_status INTEGER,
                response_time REAL,
                error_message TEXT,
                checked_at TEXT NOT NULL,
                FOREIGN KEY(component_id) REFERENCES components(id) ON DELETE CASCADE
            );
            """
        )

        count = conn.execute("SELECT COUNT(*) AS total FROM projects").fetchone()["total"]
        if count == 0:
            now = datetime.now().isoformat(timespec="seconds")
            defaults = [("Nagad", "#f57c00"), ("MBL", "#198754"), ("NexusPay", "#0d6efd")]
            conn.executemany("INSERT INTO projects(name, color, created_at) VALUES (?, ?, ?)", [(n, c, now) for n, c in defaults])




SAMPLE_IMPORT_JSON = {
    "baseUrl": "http://10.88.250.202",
    "healthPath": "/health",
    "services": [
        {"componentName": "Outbound Proxy", "port": 24424},
        {"componentName": "System Gateway", "port": 20025},
        {"componentName": "API Gateway", "port": 20005},
    ],
}


def validate_import_payload(data) -> tuple[list[dict], list[str]]:
    """Validate one configuration object or an array of configuration objects."""
    configs = data if isinstance(data, list) else [data]
    if not configs or any(not isinstance(item, dict) for item in configs):
        return [], ["Invalid JSON structure. Use a JSON object or an array of JSON objects."]

    components: list[dict] = []
    errors: list[str] = []
    names_seen: set[str] = set()

    for config_index, config in enumerate(configs, start=1):
        prefix = f"Configuration {config_index}: " if len(configs) > 1 else ""
        if not str(config.get("baseUrl", "")).strip():
            errors.append(prefix + "Missing baseUrl.")
        if not str(config.get("healthPath", "")).strip():
            errors.append(prefix + "Missing healthPath.")
        if "services" not in config:
            errors.append(prefix + "Missing services.")
            continue
        services = config.get("services")
        if not isinstance(services, list):
            errors.append(prefix + "services must be a JSON array.")
            continue
        if not services:
            errors.append(prefix + "services cannot be empty.")
            continue

        base_url = str(config.get("baseUrl", "")).rstrip("/")
        health_path = str(config.get("healthPath", ""))
        if health_path and not health_path.startswith("/"):
            health_path = "/" + health_path
        parsed = urlparse(base_url)

        for service_index, service in enumerate(services, start=1):
            service_prefix = f"{prefix}Service {service_index}: "
            if not isinstance(service, dict):
                errors.append(service_prefix + "Invalid service object.")
                continue
            name = str(service.get("componentName", "")).strip()
            if not name:
                errors.append(service_prefix + "Missing componentName.")
            port = service.get("port")
            try:
                port_number = int(port)
                if isinstance(port, bool) or port_number < 1 or port_number > 65535:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append(service_prefix + "Invalid port. Use a number from 1 to 65535.")
                continue
            key = name.casefold()
            if name and key in names_seen:
                errors.append(service_prefix + f"Duplicate component: {name}.")
            elif name:
                names_seen.add(key)
            if not name or not parsed.hostname or parsed.scheme not in {"http", "https"}:
                continue
            components.append({
                "name": name,
                "health_url": f"{parsed.scheme}://{parsed.hostname}:{port_number}{health_path}",
                "timeout": service.get("timeout", 5),
                "enabled": service.get("enabled", True),
            })
    return components, errors


def import_components_into_project(conn: sqlite3.Connection, project_id: int, components: list, duplicate_mode: str = "update") -> dict:
    added = updated = 0
    now = datetime.now().isoformat(timespec="seconds")
    for item in components:
        name = item["name"]
        existing = conn.execute(
            "SELECT id FROM components WHERE project_id=? AND lower(name)=lower(?)",
            (project_id, name),
        ).fetchone()
        values = (
            name, item["health_url"], None, None, "none", None, None, None,
            json.dumps({}), max(1, int(item.get("timeout", 5))),
            1 if item.get("enabled", True) else 0,
        )
        if existing:
            conn.execute(
                """UPDATE components SET name=?,health_url=?,info_url=?,metrics_url=?,auth_type=?,username=?,password=?,bearer_token=?,custom_headers=?,timeout=?,enabled=? WHERE id=?""",
                (*values, existing["id"]),
            )
            updated += 1
        else:
            conn.execute(
                """INSERT INTO components(project_id,name,health_url,info_url,metrics_url,auth_type,username,password,bearer_token,custom_headers,timeout,enabled,created_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (project_id, *values, now),
            )
            added += 1
    return {"added": added, "updated": updated}


def validate_url(url: str) -> bool:
    try:
        parsed = urlparse(url)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
    except Exception:
        return False


def parse_json_headers(value: str | None) -> dict:
    if not value:
        return {}
    try:
        data = json.loads(value)
        return data if isinstance(data, dict) else {}
    except json.JSONDecodeError:
        return {}


def build_request_options(component: sqlite3.Row) -> dict:
    headers = parse_json_headers(component["custom_headers"])
    auth = None
    if component["auth_type"] == "basic":
        auth = (component["username"] or "", component["password"] or "")
    elif component["auth_type"] == "bearer" and component["bearer_token"]:
        headers["Authorization"] = f"Bearer {component['bearer_token']}"
    return {"headers": headers, "auth": auth, "timeout": component["timeout"], "verify": True}


def nested_get(data: dict, *paths, default=None):
    for path in paths:
        current = data
        ok = True
        for key in path.split("."):
            if not isinstance(current, dict) or key not in current:
                ok = False
                break
            current = current[key]
        if ok:
            return current
    return default


def fetch_optional_json(url: str | None, options: dict) -> dict:
    if not url or not validate_url(url):
        return {}
    try:
        response = requests.get(url, **options)
        if response.ok:
            return response.json() if response.content else {}
    except Exception:
        pass
    return {}


def check_component(component_id: int) -> dict:
    with db_connection() as conn:
        component = conn.execute("SELECT * FROM components WHERE id = ?", (component_id,)).fetchone()
        if component is None:
            raise ValueError("Component not found")
        if not component["enabled"]:
            return {"id": component_id, "status": "UNKNOWN", "error_message": "Component is disabled"}

    checked_at = datetime.now().isoformat(timespec="seconds")
    result = {
        "status": "UNKNOWN", "http_status": None, "response_time": None,
        "version": None, "build_date": None, "git_branch": None, "git_commit": None,
        "java_version": None, "cpu_usage": None, "memory_usage": None, "disk_usage": None,
        "last_checked": checked_at, "error_message": None, "raw_response": None,
    }

    if not validate_url(component["health_url"]):
        result["error_message"] = "Invalid URL"
    else:
        options = build_request_options(component)
        started = time.perf_counter()
        try:
            response = requests.get(component["health_url"], **options)
            result["response_time"] = round((time.perf_counter() - started) * 1000, 2)
            result["http_status"] = response.status_code
            result["raw_response"] = response.text[:15000]

            payload = {}
            try:
                payload = response.json() if response.content else {}
            except ValueError:
                payload = {}
                if response.ok:
                    result["error_message"] = "Invalid JSON response"

            status_value = str(payload.get("status", "")).upper() if isinstance(payload, dict) else ""
            if response.status_code in (401, 403):
                result["status"] = "DOWN"
                result["error_message"] = "Unauthorized"
            elif response.status_code == 404:
                result["status"] = "DOWN"
                result["error_message"] = "Health endpoint not found (404)"
            elif response.status_code >= 500:
                result["status"] = "DOWN"
                result["error_message"] = f"Server error ({response.status_code})"
            elif response.ok:
                result["status"] = "UP" if status_value in {"UP", "OK", "HEALTHY", "PASS"} or not status_value else ("DOWN" if status_value in {"DOWN", "OUT_OF_SERVICE", "FAIL", "FAILED"} else "UNKNOWN")
            else:
                result["status"] = "DOWN"
                result["error_message"] = f"HTTP {response.status_code}"

            info = fetch_optional_json(component["info_url"], options)
            metrics = fetch_optional_json(component["metrics_url"], options)
            combined = {**payload, **{"info": info, "metricsData": metrics}}

            # Support common Spring Boot / custom health-response structures. Some
            # services expose build metadata directly under `application` instead
            # of a separate `/info` endpoint.
            result["version"] = nested_get(
                combined,
                "info.build.version", "info.app.version", "info.application.version",
                "info.version", "application.version", "app.version",
                "build.version", "version",
            )
            result["build_date"] = nested_get(
                combined,
                "info.build.time", "info.build.date", "info.application.buildDate",
                "application.buildDate", "application.build_date", "application.build-date",
                "build.time", "build.date", "buildDate", "build_date","build-date",
            )
            result["git_branch"] = nested_get(
                combined,
                "info.git.branch", "info.git.branch.name", "info.application.branch",
                "application.git.branch", "application.gitBranch", "application.branch",
                "git.branch", "git.branch.name", "gitBranch", "branch",
            )
            result["git_commit"] = nested_get(
                combined,
                "info.git.commit.id", "info.git.commit.id.abbrev", "info.git.commit",
                "info.application.commit", "info.application.hash",
                "application.git.commit.id", "application.git.commit",
                "application.gitCommit", "application.commit", "application.commitId",
                "application.hash", "git.commit.id", "git.commit.id.abbrev",
                "git.commit", "gitCommit", "commit", "commitId", "hash",
            )
            result["java_version"] = nested_get(
                combined,
                "info.java.version", "info.application.javaVersion",
                "application.java.version", "application.javaVersion",
                "java.version", "javaVersion",
            )
            result["cpu_usage"] = nested_get(combined, "components.cpu.details.usage", "cpuUsage")
            result["memory_usage"] = nested_get(combined, "components.memory.details.usage", "memoryUsage")
            
            ds=combined.get("diskSpace",{}) if isinstance(combined,dict) else {}
            if isinstance(ds,dict) and ds.get("free") is not None and ds.get("total") is not None:
                result["disk_usage"]=f"{ds['free']/1024**3:.2f} GB / {ds['total']/1024**3:.2f} GB"
            else:
                result["disk_usage"] = nested_get(combined, "components.diskSpace.details.free", "diskUsage")
        except requests.exceptions.Timeout:
            result["status"] = "DOWN"; result["error_message"] = "Connection timeout"
        except requests.exceptions.SSLError:
            result["status"] = "DOWN"; result["error_message"] = "SSL error"
        except requests.exceptions.ConnectionError:
            result["status"] = "DOWN"; result["error_message"] = "Connection failed"
        except requests.exceptions.RequestException as exc:
            result["status"] = "DOWN"; result["error_message"] = str(exc)
        except Exception as exc:
            result["status"] = "UNKNOWN"; result["error_message"] = f"Unexpected error: {exc}"

    with db_connection() as conn:
        conn.execute(
            """UPDATE components SET status=?, http_status=?, response_time=?, version=?, build_date=?, git_branch=?, git_commit=?, java_version=?, cpu_usage=?, memory_usage=?, disk_usage=?, last_checked=?, error_message=?, raw_response=? WHERE id=?""",
            (result["status"], result["http_status"], result["response_time"], result["version"], result["build_date"], result["git_branch"], result["git_commit"], result["java_version"], result["cpu_usage"], result["memory_usage"], result["disk_usage"], checked_at, result["error_message"], result["raw_response"], component_id),
        )
        conn.execute("INSERT INTO health_history(component_id,status,http_status,response_time,error_message,checked_at) VALUES(?,?,?,?,?,?)",
                     (component_id, result["status"], result["http_status"], result["response_time"], result["error_message"], checked_at))
    logger.info("Health check | component=%s | status=%s | http=%s | time_ms=%s | error=%s", component["name"], result["status"], result["http_status"], result["response_time"], result["error_message"])
    return {"id": component_id, **result}


@app.before_request
def log_request():
    logger.info("Request | %s %s", request.method, request.path)


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/dashboard")
def dashboard_data():
    with db_connection() as conn:
        projects = [dict(r) for r in conn.execute("SELECT * FROM projects ORDER BY name").fetchall()]
        components = [dict(r) for r in conn.execute("SELECT c.*, p.name AS project_name, p.color AS project_color FROM components c JOIN projects p ON p.id=c.project_id ORDER BY p.name,c.name").fetchall()]
    summary = {
        "projects": len(projects), "components": len(components),
        "up": sum(1 for c in components if c["status"] == "UP"),
        "down": sum(1 for c in components if c["status"] == "DOWN"),
        "unknown": sum(1 for c in components if c["status"] == "UNKNOWN"),
    }
    times = [c["response_time"] for c in components if c["response_time"] is not None]
    summary["average_response_time"] = round(sum(times) / len(times), 2) if times else 0
    return jsonify({"projects": projects, "components": components, "summary": summary})


@app.post("/api/projects")
def create_project():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    color = str(data.get("color", "#0d6efd")).strip()
    if not name:
        return jsonify({"error": "Project name is required"}), 400
    try:
        with db_connection() as conn:
            cur = conn.execute("INSERT INTO projects(name,color,created_at) VALUES(?,?,?)", (name, color, datetime.now().isoformat(timespec="seconds")))
        return jsonify({"id": cur.lastrowid, "message": "Project created"}), 201
    except sqlite3.IntegrityError:
        return jsonify({"error": "Project name already exists"}), 409


@app.put("/api/projects/<int:project_id>")
def update_project(project_id):
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip(); color = str(data.get("color", "#0d6efd")).strip()
    if not name: return jsonify({"error": "Project name is required"}), 400
    try:
        with db_connection() as conn:
            conn.execute("UPDATE projects SET name=?, color=? WHERE id=?", (name, color, project_id))
        return jsonify({"message": "Project updated"})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Project name already exists"}), 409


@app.delete("/api/projects/<int:project_id>")
def delete_project(project_id):
    with db_connection() as conn:
        project = conn.execute("SELECT id FROM projects WHERE id=?", (project_id,)).fetchone()
        if project is None:
            return jsonify({"error": "Project not found"}), 404
        component_ids = [row["id"] for row in conn.execute("SELECT id FROM components WHERE project_id=?", (project_id,)).fetchall()]
        if component_ids:
            placeholders = ",".join("?" for _ in component_ids)
            conn.execute(f"DELETE FROM health_history WHERE component_id IN ({placeholders})", component_ids)
        conn.execute("DELETE FROM components WHERE project_id=?", (project_id,))
        conn.execute("DELETE FROM projects WHERE id=?", (project_id,))
    return jsonify({"message": "Project and its components deleted"})


@app.post("/api/components")
def create_component():
    data = request.get_json(silent=True) or {}
    required = ["project_id", "name", "health_url"]
    if any(not data.get(k) for k in required): return jsonify({"error": "Project, component name and health URL are required"}), 400
    if not validate_url(data["health_url"]): return jsonify({"error": "Invalid health URL"}), 400
    with db_connection() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?", (int(data["project_id"]),)).fetchone():
            return jsonify({"error": "Selected project does not exist"}), 404
        cur = conn.execute(
            """INSERT INTO components(project_id,name,health_url,info_url,metrics_url,auth_type,username,password,bearer_token,custom_headers,timeout,enabled,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (data["project_id"], data["name"].strip(), data["health_url"].strip(), data.get("info_url") or None, data.get("metrics_url") or None,
             data.get("auth_type", "none"), data.get("username"), data.get("password"), data.get("bearer_token"), json.dumps(data.get("custom_headers") or {}),
             int(data.get("timeout", 5)), 1 if data.get("enabled", True) else 0, datetime.now().isoformat(timespec="seconds")),
        )
    return jsonify({"id": cur.lastrowid, "message": "Component created"}), 201


@app.put("/api/components/<int:component_id>")
def update_component(component_id):
    data = request.get_json(silent=True) or {}
    if not data.get("project_id") or not data.get("name") or not data.get("health_url"): return jsonify({"error": "Required fields are missing"}), 400
    if not validate_url(data["health_url"]): return jsonify({"error": "Invalid health URL"}), 400
    with db_connection() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?", (int(data["project_id"]),)).fetchone():
            return jsonify({"error": "Selected project does not exist"}), 404
        if not conn.execute("SELECT 1 FROM components WHERE id=?", (component_id,)).fetchone():
            return jsonify({"error": "Component not found"}), 404
        conn.execute(
            """UPDATE components SET project_id=?,name=?,health_url=?,info_url=?,metrics_url=?,auth_type=?,username=?,password=?,bearer_token=?,custom_headers=?,timeout=?,enabled=? WHERE id=?""",
            (data["project_id"], data["name"].strip(), data["health_url"].strip(), data.get("info_url") or None, data.get("metrics_url") or None,
             data.get("auth_type", "none"), data.get("username"), data.get("password"), data.get("bearer_token"), json.dumps(data.get("custom_headers") or {}),
             int(data.get("timeout", 5)), 1 if data.get("enabled", True) else 0, component_id),
        )
    return jsonify({"message": "Component updated"})


@app.delete("/api/components/<int:component_id>")
def delete_component(component_id):
    with db_connection() as conn:
        if not conn.execute("SELECT 1 FROM components WHERE id=?", (component_id,)).fetchone():
            return jsonify({"error": "Component not found"}), 404
        conn.execute("DELETE FROM health_history WHERE component_id=?", (component_id,))
        conn.execute("DELETE FROM components WHERE id=?", (component_id,))
    return jsonify({"message": "Component deleted"})


@app.post("/api/check/component/<int:component_id>")
def check_one(component_id):
    try: return jsonify(check_component(component_id))
    except ValueError as exc: return jsonify({"error": str(exc)}), 404


@app.post("/api/check/project/<int:project_id>")
def check_project(project_id):
    with db_connection() as conn:
        ids = [r["id"] for r in conn.execute("SELECT id FROM components WHERE project_id=? AND enabled=1", (project_id,)).fetchall()]
    return jsonify({"results": [check_component(i) for i in ids]})


@app.post("/api/check/all")
def check_all():
    with db_connection() as conn:
        ids = [r["id"] for r in conn.execute("SELECT id FROM components WHERE enabled=1").fetchall()]
    return jsonify({"results": [check_component(i) for i in ids]})


@app.get("/api/import/sample")
def download_import_sample():
    buffer = BytesIO(json.dumps(SAMPLE_IMPORT_JSON, indent=2).encode("utf-8"))
    return send_file(buffer, mimetype="application/json", as_attachment=True, download_name="enterprise_health_dashboard_sample.json")


@app.post("/api/import/project/<int:project_id>")
def import_project_json(project_id):
    data = request.get_json(silent=True)
    if data is None:
        return jsonify({"error": "Invalid JSON. Check commas, quotation marks, and brackets."}), 400
    components, errors = validate_import_payload(data)
    if errors:
        return jsonify({"error": " ".join(errors), "errors": errors}), 400
    import_mode = request.args.get("import_mode", "merge")
    if import_mode not in {"merge", "replace"}:
        return jsonify({"error": "Invalid import mode."}), 400
    with db_connection() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone():
            return jsonify({"error": "Project not found"}), 404
        if import_mode == "replace":
            conn.execute("DELETE FROM components WHERE project_id=?", (project_id,))
        result = import_components_into_project(conn, project_id, components)
    return jsonify({"message": f"Import completed: {result['added']} added, {result['updated']} updated", **result})


@app.post("/api/import/project/<int:project_id>/files")
def import_project_json_files(project_id):
    files = request.files.getlist("files")
    import_mode = request.form.get("import_mode", "merge")
    if not files:
        return jsonify({"error": "Please select at least one JSON file."}), 400
    if import_mode not in {"merge", "replace"}:
        return jsonify({"error": "Invalid import mode."}), 400

    parsed_components: list[dict] = []
    all_errors: list[str] = []
    names_seen: set[str] = set()
    for uploaded in files:
        filename = uploaded.filename or "unnamed.json"
        if not filename.lower().endswith(".json"):
            all_errors.append(f"{filename}: Unsupported file type. Only .json files are allowed.")
            continue
        try:
            data = json.load(uploaded.stream)
        except json.JSONDecodeError as exc:
            all_errors.append(f"{filename}: Invalid JSON near line {exc.lineno}, column {exc.colno}.")
            continue
        except UnicodeDecodeError:
            all_errors.append(f"{filename}: File encoding must be UTF-8.")
            continue
        components, errors = validate_import_payload(data)
        all_errors.extend(f"{filename}: {error}" for error in errors)
        for component in components:
            key = component["name"].casefold()
            if key in names_seen:
                all_errors.append(f"{filename}: Duplicate component: {component['name']}.")
            else:
                names_seen.add(key)
                parsed_components.append(component)

    if all_errors:
        return jsonify({"error": " ".join(all_errors), "errors": all_errors}), 400
    with db_connection() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone():
            return jsonify({"error": "Project not found"}), 404
        if import_mode == "replace":
            conn.execute("DELETE FROM components WHERE project_id=?", (project_id,))
        result = import_components_into_project(conn, project_id, parsed_components)
    return jsonify({"message": f"{len(files)} file(s) imported: {result['added']} added, {result['updated']} updated", **result})


@app.get("/api/export/project/<int:project_id>")
def export_project_json(project_id):
    with db_connection() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
        if not project: return jsonify({"error": "Project not found"}), 404
        comps = [dict(r) for r in conn.execute("SELECT name,health_url,info_url,metrics_url,auth_type,username,bearer_token,custom_headers,timeout,enabled FROM components WHERE project_id=? ORDER BY name", (project_id,)).fetchall()]
    payload = {"project": {"name": project["name"], "color": project["color"]}, "components": comps}
    buffer = BytesIO(json.dumps(payload, indent=2).encode("utf-8"))
    return send_file(buffer, mimetype="application/json", as_attachment=True, download_name=f"{project['name']}_config.json")


@app.get("/api/export/json")
def export_all_json():
    data = dashboard_data().get_json()
    buffer = BytesIO(json.dumps(data, indent=2, default=str).encode("utf-8"))
    return send_file(buffer, mimetype="application/json", as_attachment=True, download_name="enterprise_health_dashboard.json")


@app.get("/api/export/excel")
def export_excel():
    with db_connection() as conn:
        rows = conn.execute("SELECT p.name project,c.* FROM components c JOIN projects p ON p.id=c.project_id ORDER BY p.name,c.name").fetchall()
    wb = Workbook(); ws = wb.active; ws.title = "Health Dashboard"
    headers = ["Project","Component","Status","HTTP Status","Response Time (ms)","Version","Build Date","Git Branch","Git Commit","Java Version","CPU","Memory","Disk","Last Checked","Error"]
    ws.append(headers)
    for r in rows: ws.append([r["project"],r["name"],r["status"],r["http_status"],r["response_time"],r["version"],r["build_date"],r["git_branch"],r["git_commit"],r["java_version"],r["cpu_usage"],r["memory_usage"],r["disk_usage"],r["last_checked"],r["error_message"]])
    for col in ws.columns:
        width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
        ws.column_dimensions[col[0].column_letter].width = width
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="enterprise_health_dashboard.xlsx")


@app.get("/api/export/pdf")
def export_pdf():
    with db_connection() as conn:
        rows = conn.execute("SELECT p.name project,c.name,c.status,c.http_status,c.response_time,c.last_checked FROM components c JOIN projects p ON p.id=c.project_id ORDER BY p.name,c.name").fetchall()
    output = BytesIO(); pdf = canvas.Canvas(output, pagesize=A4); width, height = A4
    y = height - 50; pdf.setFont("Helvetica-Bold", 15); pdf.drawString(40, y, "Enterprise Health Dashboard"); y -= 30
    pdf.setFont("Helvetica", 8)
    for r in rows:
        line = f"{r['project']} | {r['name']} | {r['status']} | HTTP {r['http_status'] or '-'} | {r['response_time'] or '-'} ms | {r['last_checked'] or '-'}"
        pdf.drawString(40, y, line[:120]); y -= 14
        if y < 40: pdf.showPage(); pdf.setFont("Helvetica", 8); y = height - 40
    pdf.save(); output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name="enterprise_health_dashboard.pdf")


@app.get("/api/history/<int:component_id>")
def component_history(component_id):
    with db_connection() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM health_history WHERE component_id=? ORDER BY id DESC LIMIT 50", (component_id,)).fetchall()]
    return jsonify(rows)


@app.errorhandler(Exception)
def handle_error(exc):
    logger.exception("Unhandled error: %s", exc)
    return jsonify({"error": "Unexpected application error"}), 500


if __name__ == "__main__":
    init_db()
    # Port 5001 so it runs alongside the AutoRCA portal (which uses 5000).
    app.run(host="127.0.0.1", port=5001, debug=False)
